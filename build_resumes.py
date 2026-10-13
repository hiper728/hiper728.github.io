#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generate Harvard/Silicon Valley ATS-Compliant PDF & Word Resumes.
Output files:
- assets/Leonard_Johnson_Resume.pdf
- assets/Leonard_Johnson_Resume.docx
"""

import os
from pathlib import Path

# --- PDF GENERATION WITH REPORTLAB ---
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT

OUT_DIR = Path(r"C:\Users\hiper\.gemini\antigravity\scratch\hiper728_portfolio\assets")
OUT_DIR.mkdir(parents=True, exist_ok=True)
PDF_PATH = OUT_DIR / "Leonard_Johnson_Resume.pdf"
DOCX_PATH = OUT_DIR / "Leonard_Johnson_Resume.docx"

def generate_pdf():
    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36,
        title="Leonard Johnson - Senior Software & Automation Engineer Resume",
        author="Leonard Johnson",
    )

    styles = getSampleStyleSheet()
    
    # Custom styles
    name_style = ParagraphStyle(
        "DocName",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#0f172a"),
    )
    
    contact_style = ParagraphStyle(
        "DocContact",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#334155"),
    )
    
    section_title_style = ParagraphStyle(
        "SectionTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#0f172a"),
        spaceBefore=8,
        spaceAfter=2,
    )
    
    job_header_left = ParagraphStyle(
        "JobHeaderLeft",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#1e293b"),
    )
    
    job_header_right = ParagraphStyle(
        "JobHeaderRight",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13,
        alignment=TA_RIGHT,
        textColor=colors.HexColor("#475569"),
    )
    
    body_style = ParagraphStyle(
        "BodyTextCustom",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12.5,
        textColor=colors.HexColor("#334155"),
    )

    bullet_style = ParagraphStyle(
        "BulletTextCustom",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12.5,
        leftIndent=12,
        firstLineIndent=-8,
        textColor=colors.HexColor("#1e293b"),
        spaceAfter=2.5,
    )

    story = []

    # Header
    story.append(Paragraph("LEONARD JOHNSON", name_style))
    story.append(Spacer(1, 3))
    story.append(Paragraph("Senior Software & Systems Automation Engineer", ParagraphStyle("TitleSub", parent=contact_style, fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=colors.HexColor("#2563eb"))))
    story.append(Spacer(1, 4))
    
    contacts = "a.t.87.14.912@gmail.com &bull; Telegram: @js122b &bull; github.com/hiper728 &bull; hiper728.github.io"
    story.append(Paragraph(contacts, contact_style))
    story.append(Spacer(1, 8))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cbd5e1"), spaceBefore=2, spaceAfter=6))

    # Professional Summary
    story.append(Paragraph("PROFESSIONAL SUMMARY", section_title_style))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#94a3b8"), spaceBefore=1, spaceAfter=5))
    summary_text = (
        "Senior Software Engineer with 6+ years of specialized experience architecting high-throughput asynchronous "
        "backends, resilient automation systems, and high-performance analytical pipelines. Proven track record of "
        "building production-grade bot ecosystems, persistent browser automation engines, and containerized zero-downtime "
        "Linux deployment workflows. Dedicated to writing clean, maintainable code with strict type-safety and observable architectures."
    )
    story.append(Paragraph(summary_text, body_style))
    story.append(Spacer(1, 6))

    # Technical Skills
    story.append(Paragraph("TECHNICAL EXPERTISE", section_title_style))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#94a3b8"), spaceBefore=1, spaceAfter=5))
    skills = [
        "<b>Languages & Runtimes:</b> Python 3.13 (Asyncio, CPython Internals), SQL, Bash/Shell, Go, TypeScript/JavaScript",
        "<b>Frameworks & Automation:</b> FastAPI, aiogram 3.x, discord.py, Playwright, Puppeteer, SQLAlchemy, Pydantic v2",
        "<b>Data & Databases:</b> PostgreSQL, DuckDB, SQLite (WAL mode), Redis, In-Memory Caching, Schema Migrations",
        "<b>Infrastructure & DevOps:</b> Linux (Debian, Ubuntu), Docker, Docker Compose, Nginx Reverse Proxy, Systemd, GitHub Actions CI/CD",
        "<b>Core Competencies:</b> High-Throughput Message Dispatch, Web Scraping & Stealth Automation, Database Optimization, Microservices, API Security",
    ]
    for s in skills:
        story.append(Paragraph(f"&bull; {s}", bullet_style))
    story.append(Spacer(1, 6))

    # Experience
    story.append(Paragraph("WORK EXPERIENCE", section_title_style))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#94a3b8"), spaceBefore=1, spaceAfter=5))

    # Role 1
    t1 = Table([[Paragraph("Lead Systems & Automation Engineer", job_header_left), Paragraph("2022 &ndash; Present", job_header_right)]], colWidths=[380, 160])
    t1.setStyle(TableStyle([('VALIGN', (0,0), (-1,-1), 'TOP'), ('LEFTPADDING', (0,0), (-1,-1), 0), ('RIGHTPADDING', (0,0), (-1,-1), 0)]))
    story.append(t1)
    story.append(Paragraph("<i>Distributed Systems & Infrastructure Engineering</i>", body_style))
    story.append(Spacer(1, 3))
    r1_points = [
        "Architected an asynchronous messaging engine coordinating 20+ worker nodes, handling 450+ msg/sec peak throughput with sub-40ms latency and 99.98% reliability.",
        "Engineered a resilient headless browser scraping engine using Playwright and CDP, monitoring 50,000+ endpoints daily while reducing anti-bot blocking by 94% through session persistence and header randomization.",
        "Constructed in-process analytical data pipelines using DuckDB and Polars, accelerating analytical batch query processing 10x over relational databases.",
        "Containerized production microservices with optimized multi-stage Dockerfiles and established automated GitHub Actions CI/CD deployment to Linux servers behind Nginx.",
    ]
    for p in r1_points:
        story.append(Paragraph(f"&bull; {p}", bullet_style))
    story.append(Spacer(1, 5))

    # Role 2
    t2 = Table([[Paragraph("Senior Backend Engineer", job_header_left), Paragraph("2019 &ndash; 2022", job_header_right)]], colWidths=[380, 160])
    t2.setStyle(TableStyle([('VALIGN', (0,0), (-1,-1), 'TOP'), ('LEFTPADDING', (0,0), (-1,-1), 0), ('RIGHTPADDING', (0,0), (-1,-1), 0)]))
    story.append(t2)
    story.append(Paragraph("<i>High-Scale APIs & Automation Platforms</i>", body_style))
    story.append(Spacer(1, 3))
    r2_points = [
        "Built robust REST API services and asynchronous background workers using FastAPI, Pydantic, and PostgreSQL serving over 1M monthly requests.",
        "Implemented automated database backup, migration (Alembic), and point-in-time recovery workflows across distributed SQLite and PostgreSQL databases.",
        "Spearheaded automated testing strategies with pytest achieving 92% test coverage across core messaging and financial transaction services.",
    ]
    for p in r2_points:
        story.append(Paragraph(f"&bull; {p}", bullet_style))
    story.append(Spacer(1, 6))

    # Key Projects
    story.append(Paragraph("FEATURED ENGINEERING PROJECTS", section_title_style))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#94a3b8"), spaceBefore=1, spaceAfter=5))
    projects = [
        "<b>OmniBot Distributed Platform:</b> Production-ready Telegram & Discord bot engine featuring state-machine persistence, admin rate-limiting, and SQLite WAL transactional logging.",
        "<b>Sentinel Web Scraper:</b> Headless browser automation system with visual DOM diffing, persistent session caching, and webhook alerting.",
        "<b>DuckDB High-Velocity Analytics:</b> Local query engine enabling sub-second analytical SQL aggregation over multi-gigabyte CSV and Parquet files without database overhead.",
    ]
    for prj in projects:
        story.append(Paragraph(f"&bull; {prj}", bullet_style))
    story.append(Spacer(1, 6))

    # Education
    story.append(Paragraph("EDUCATION & CERTIFICATIONS", section_title_style))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#94a3b8"), spaceBefore=1, spaceAfter=5))
    story.append(Paragraph("<b>B.Sc. in Computer Science & Software Engineering</b>", job_header_left))
    story.append(Paragraph("&bull; Advanced Linux System Administration & Container Orchestration", bullet_style))
    story.append(Paragraph("&bull; Asynchronous Concurrency and Distributed Systems Architecture", bullet_style))

    doc.build(story)
    print(f"Generated PDF: {PDF_PATH}")


# --- DOCX GENERATION ---
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_margins(cell, top=50, bottom=50, left=50, right=50):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def generate_docx():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.5)
        s.bottom_margin = Inches(0.5)
        s.left_margin = Inches(0.6)
        s.right_margin = Inches(0.6)

    # Title
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_name = p_name.add_run("LEONARD JOHNSON\n")
    r_name.font.name = "Calibri"
    r_name.font.size = Pt(20)
    r_name.bold = True
    r_name.font.color.rgb = RGBColor(15, 23, 42)

    r_title = p_name.add_run("Senior Software & Systems Automation Engineer\n")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(11)
    r_title.bold = True
    r_title.font.color.rgb = RGBColor(37, 99, 235)

    r_contact = p_name.add_run("a.t.87.14.912@gmail.com  |  Telegram: @js122b  |  github.com/hiper728  |  hiper728.github.io")
    r_contact.font.name = "Calibri"
    r_contact.font.size = Pt(9.5)
    r_contact.font.color.rgb = RGBColor(71, 85, 105)

    def add_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.bold = True
        r.font.color.rgb = RGBColor(15, 23, 42)
        
        # Border bottom via xml
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), '94A3B8')
        pBdr.append(bottom)
        pPr.append(pBdr)

    # Summary
    add_section_header("PROFESSIONAL SUMMARY")
    p_sum = doc.add_paragraph(
        "Senior Software Engineer with 6+ years of specialized experience architecting high-throughput asynchronous "
        "backends, resilient automation systems, and high-performance analytical pipelines. Proven track record of "
        "building production-grade bot ecosystems, persistent browser automation engines, and containerized zero-downtime "
        "Linux deployment workflows. Dedicated to writing clean, maintainable code with strict type-safety and observable architectures."
    )
    p_sum.paragraph_format.space_after = Pt(4)
    p_sum.style.font.name = "Calibri"
    p_sum.style.font.size = Pt(9.5)

    # Technical Expertise
    add_section_header("TECHNICAL EXPERTISE")
    skills = [
        ("Languages & Runtimes: ", "Python 3.13 (Asyncio, CPython Internals), SQL, Bash/Shell, Go, TypeScript/JavaScript"),
        ("Frameworks & Automation: ", "FastAPI, aiogram 3.x, discord.py, Playwright, Puppeteer, SQLAlchemy, Pydantic v2"),
        ("Data & Databases: ", "PostgreSQL, DuckDB, SQLite (WAL mode), Redis, In-Memory Caching, Schema Migrations"),
        ("Infrastructure & DevOps: ", "Linux (Debian, Ubuntu), Docker, Docker Compose, Nginx Reverse Proxy, Systemd, GitHub Actions CI/CD"),
        ("Core Competencies: ", "High-Throughput Message Dispatch, Web Scraping & Stealth Automation, Database Optimization, Microservices, API Security"),
    ]
    for k, v in skills:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        r_k = p.add_run(k)
        r_k.bold = True
        r_k.font.name = "Calibri"
        r_k.font.size = Pt(9.5)
        r_v = p.add_run(v)
        r_v.font.name = "Calibri"
        r_v.font.size = Pt(9.5)

    # Experience
    add_section_header("WORK EXPERIENCE")
    
    # Job 1
    p_j1 = doc.add_paragraph()
    p_j1.paragraph_format.space_after = Pt(1)
    r1 = p_j1.add_run("Lead Systems & Automation Engineer  |  2022 – Present")
    r1.bold = True
    r1.font.name = "Calibri"
    r1.font.size = Pt(10)
    
    j1_bullets = [
        "Architected an asynchronous messaging engine coordinating 20+ worker nodes, handling 450+ msg/sec peak throughput with sub-40ms latency and 99.98% reliability.",
        "Engineered a resilient headless browser scraping engine using Playwright and CDP, monitoring 50,000+ endpoints daily while reducing anti-bot blocking by 94% through session persistence and header randomization.",
        "Constructed in-process analytical data pipelines using DuckDB and Polars, accelerating analytical batch query processing 10x over relational databases.",
        "Containerized production microservices with optimized multi-stage Dockerfiles and established automated GitHub Actions CI/CD deployment to Linux servers behind Nginx.",
    ]
    for b in j1_bullets:
        p = doc.add_paragraph(b, style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        p.style.font.name = "Calibri"
        p.style.font.size = Pt(9.5)

    # Job 2
    p_j2 = doc.add_paragraph()
    p_j2.paragraph_format.space_before = Pt(4)
    p_j2.paragraph_format.space_after = Pt(1)
    r2 = p_j2.add_run("Senior Backend Engineer  |  2019 – 2022")
    r2.bold = True
    r2.font.name = "Calibri"
    r2.font.size = Pt(10)
    
    j2_bullets = [
        "Built robust REST API services and asynchronous background workers using FastAPI, Pydantic, and PostgreSQL serving over 1M monthly requests.",
        "Implemented automated database backup, migration (Alembic), and point-in-time recovery workflows across distributed SQLite and PostgreSQL databases.",
        "Spearheaded automated testing strategies with pytest achieving 92% test coverage across core messaging and financial transaction services.",
    ]
    for b in j2_bullets:
        p = doc.add_paragraph(b, style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        p.style.font.name = "Calibri"
        p.style.font.size = Pt(9.5)

    # Projects
    add_section_header("FEATURED ENGINEERING PROJECTS")
    prjs = [
        ("OmniBot Distributed Platform: ", "Production-ready Telegram & Discord bot engine featuring state-machine persistence, admin rate-limiting, and SQLite WAL transactional logging."),
        ("Sentinel Web Scraper: ", "Headless browser automation system with visual DOM diffing, persistent session caching, and webhook alerting."),
        ("DuckDB High-Velocity Analytics: ", "Local query engine enabling sub-second analytical SQL aggregation over multi-gigabyte CSV and Parquet files without database overhead."),
    ]
    for k, v in prjs:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        r_k = p.add_run(k)
        r_k.bold = True
        r_k.font.name = "Calibri"
        r_k.font.size = Pt(9.5)
        r_v = p.add_run(v)
        r_v.font.name = "Calibri"
        r_v.font.size = Pt(9.5)

    # Education
    add_section_header("EDUCATION & CERTIFICATIONS")
    p_edu = doc.add_paragraph()
    r_edu = p_edu.add_run("B.Sc. in Computer Science & Software Engineering\n")
    r_edu.bold = True
    r_edu.font.name = "Calibri"
    r_edu.font.size = Pt(9.5)
    p_c1 = doc.add_paragraph("Advanced Linux System Administration & Container Orchestration", style='List Bullet')
    p_c1.paragraph_format.space_after = Pt(1)
    p_c2 = doc.add_paragraph("Asynchronous Concurrency and Distributed Systems Architecture", style='List Bullet')
    p_c2.paragraph_format.space_after = Pt(1)

    doc.save(str(DOCX_PATH))
    print(f"Generated DOCX: {DOCX_PATH}")

if __name__ == "__main__":
    generate_pdf()
    generate_docx()
